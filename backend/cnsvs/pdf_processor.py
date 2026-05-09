"""
PDF Processing Service
Main service that combines limbic and cognitive processing.
"""
import os
import tempfile
from io import BytesIO
from datetime import datetime
import uuid
from typing import Tuple, Optional, Union
import pandas as pd
import pdfplumber
from docx import Document

from .limbic_service import TestIdentifier, read_pdf
from .cognitive_service import (
    extract_cognitive_data,
    extract_standard_scores_from_table,
    extract_test_standard_scores_from_tables,
    docx_with_cognitive_section_and_scores,
)
from .cognitive_service import CANONICAL_TESTS
from .plotting_service import extract_longitudinal_table, plot_longitudinal_table
from .html_service import (
    create_html_document,
    create_limbic_html_section,
    create_cognitive_html_section,
    create_longitudinal_html_section,
    create_flagged_findings_summary,
    html_page_break,
)


def combine_reports(pdf_path: str, output_path: Optional[str] = None) -> Union[BytesIO, str]:
    """
    Combine limbic and cognitive reports into a single Word document.
    
    Args:
        pdf_path: Path to input PDF file
        output_path: Optional path to save output file
        
    Returns:
        BytesIO object containing the DOCX file, or file path if output_path provided
    """
    identifier = TestIdentifier()
    pdf_text = read_pdf(pdf_path)
    identified_tests = identifier.identify_tests(pdf_text)
    test_sections = identifier.get_test_sections(pdf_text)
    
    doc = Document()
    doc = identifier.create_word_document(identified_tests, test_sections, doc=doc)
    doc.add_page_break()
    
    # Extract cognitive data
    cognitive_tables, cognitive_test_data = extract_cognitive_data(pdf_path)
    df1 = cognitive_tables.get('upper_table')
    df2 = cognitive_tables.get('lower_table')
    
    # Extract tables for standard scores
    with pdfplumber.open(pdf_path) as pdf:
        tables = []
        for page in pdf.pages[:3]:
            extracted_tables = page.extract_tables()
            tables.extend(extracted_tables)
    
    standard_scores = extract_standard_scores_from_table(tables[1]) if len(tables) > 1 else []
    test_standard_scores = extract_test_standard_scores_from_tables(tables)
    
    # Combine dataframes
    if df1 is not None and df2 is not None:
        combined_df = pd.concat([df1, df2], ignore_index=True)
    elif df1 is not None:
        combined_df = df1
    elif df2 is not None:
        combined_df = df2
    else:
        combined_df = pd.DataFrame()
    
    # Add cognitive section
    doc = docx_with_cognitive_section_and_scores(
        combined_df, cognitive_test_data, standard_scores,
        test_standard_scores, df1, tables, doc=doc
    )
    
    # Add longitudinal data if available
    long_df = extract_longitudinal_table(pdf_path)
    if long_df is not None and not long_df.empty:
        img_bytes = plot_longitudinal_table(long_df)
        doc.add_page_break()
        doc.add_picture(
            img_bytes,
            width=doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        )
        doc.add_paragraph('Longitudinal Cognitive Domain Scores (Line Plot)')
        
        # Add additional pairwise plots
        if hasattr(img_bytes, 'pair_plot_bytes'):
            pair_plot_bytes = img_bytes.pair_plot_bytes
            pair_plot_labels = getattr(img_bytes, 'pair_plot_labels', None)
            for i, pair_img in enumerate(pair_plot_bytes):
                doc.add_page_break()
                if pair_plot_labels and i < len(pair_plot_labels):
                    doc.add_paragraph(pair_plot_labels[i])
                doc.add_picture(
                    pair_img,
                    width=doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
                )
    
    if output_path:
        doc.save(output_path)
        return output_path
    else:
        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)
        return bio


def combine_reports_html(pdf_path: str, output_path: Optional[str] = None) -> Union[BytesIO, str]:
    """
    Combine limbic and cognitive reports into a single HTML document.
    
    Args:
        pdf_path: Path to input PDF file
        output_path: Optional path to save output file
        
    Returns:
        BytesIO object containing the HTML file, or file path if output_path provided
    """
    identifier = TestIdentifier()
    pdf_text = read_pdf(pdf_path)
    identified_tests = identifier.identify_tests(pdf_text)
    test_sections = identifier.get_test_sections(pdf_text)
    
    # Use original test sections directly
    sections_to_use = test_sections
    
    # Create limbic HTML section
    limbic_html = create_limbic_html_section(
        identified_tests,
        sections_to_use,
        identifier.test_patterns
    )
    
    # Extract cognitive data
    cognitive_tables, cognitive_test_data = extract_cognitive_data(pdf_path)
    df1 = cognitive_tables.get('upper_table')
    df2 = cognitive_tables.get('lower_table')
    
    # Extract tables for standard scores
    with pdfplumber.open(pdf_path) as pdf:
        tables = []
        for page in pdf.pages[:3]:
            extracted_tables = page.extract_tables()
            tables.extend(extracted_tables)
    
    standard_scores = extract_standard_scores_from_table(tables[1]) if len(tables) > 1 else []
    test_standard_scores = extract_test_standard_scores_from_tables(tables)
    
    # Collect flagged findings
    flagged_findings = []
    
    # Create cognitive HTML section (this will populate flagged_findings)
    cognitive_html = create_cognitive_html_section(
        df1,
        cognitive_test_data,
        standard_scores,
        test_standard_scores,
        CANONICAL_TESTS,
        flagged_findings=flagged_findings
    )
    
    # Combine HTML sections
    html_content = limbic_html + "\n" + cognitive_html
    
    # Add longitudinal data if available
    long_df = extract_longitudinal_table(pdf_path)
    if long_df is not None and not long_df.empty:
        img_bytes = plot_longitudinal_table(long_df)
        pair_plot_bytes = getattr(img_bytes, 'pair_plot_bytes', [])
        pair_plot_labels = getattr(img_bytes, 'pair_plot_labels', [])
        
        longitudinal_html = create_longitudinal_html_section(
            img_bytes,
            pair_plot_bytes,
            pair_plot_labels
        )
        html_content += "\n" + longitudinal_html
    
    # Add flagged findings summary at the end if there are any
    if flagged_findings:
        flagged_summary = create_flagged_findings_summary(flagged_findings)
        html_content += "\n" + flagged_summary
    
    # Create complete HTML document
    full_html = create_html_document(html_content)
    
    if output_path:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(full_html)
        return output_path
    else:
        bio = BytesIO()
        bio.write(full_html.encode('utf-8'))
        bio.seek(0)
        return bio


def generate_unique_filename(original_filename: str, extension: str = ".docx") -> str:
    """
    Generate a unique filename with timestamp and UUID.
    
    Args:
        original_filename: Original filename
        extension: File extension (default: .docx)
        
    Returns:
        Unique filename string
    """
    file_ext = os.path.splitext(original_filename)[1]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_id = str(uuid.uuid4())[:8]
    base_name = file_ext.replace('.pdf', '')
    return f"{timestamp}_{unique_id}{base_name}{extension}"

