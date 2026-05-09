"""
Cognitive Data Extraction Service
Handles extraction of cognitive test data from PDFs.
"""
import pdfplumber
import pandas as pd
from docx import Document
from docx.shared import Pt
from typing import Dict, List, Tuple, Optional

# Specific Test Scores
SPECIFIC_TESTS = [
    "Verbal Memory Test (VBM)",
    "Visual Memory Test (VSM)",
    "Finger Tapping Test (FTT)",
    "Symbol Digit Coding (SDC)",
    "Stroop Test (ST)",
    "Shifting Attention Test (SAT)",
    "Continuous Performance Test (CPT)",
    "Perception Of Emotions Test (POET)",
    "Positive Emotions",
    "Negative Emotions",
    "Reasoning Test (RT)",
]

# Four Part Continuous Performance Test
FOUR_PART_TESTS = [
    "Part 1",
    "Part 2",
    "Part 3",
    "Part 4",
]

CANONICAL_TESTS = SPECIFIC_TESTS + ["Four Part Continuous Performance Test"] + FOUR_PART_TESTS


def clean_percentile(value) -> Optional[int]:
    """Clean percentile value to integer."""
    try:
        return int(''.join(filter(str.isdigit, str(value))))
    except:
        return None


def clean_score(value) -> Optional[int]:
    """Clean score value to integer."""
    try:
        return int(''.join(filter(str.isdigit, str(value))))
    except:
        return None


def extract_cognitive_data(pdf_path: str) -> Tuple[Dict[str, pd.DataFrame], Dict]:
    """
    Extract cognitive data from PDF.
    
    Args:
        pdf_path: Path to PDF file
        
    Returns:
        Tuple of (cognitive_tables dict, cognitive_test_data dict)
    """
    with pdfplumber.open(pdf_path) as pdf:
        tables = []
        for page in pdf.pages[:3]:  # Only first 3 pages for cognitive
            extracted_tables = page.extract_tables()
            tables.extend(extracted_tables)

    # Upper Table (Patient Profile)
    df1 = None
    if len(tables) > 1:
        table_data_1 = tables[1]
        cleaned_columns_1 = []
        for i, column_name in enumerate(table_data_1[0]):
            if column_name is None:
                cleaned_columns_1.append(f"Unnamed_{i}")
            elif table_data_1[0].count(column_name) > 1:
                cleaned_columns_1.append(f"{column_name}_{i}")
            else:
                cleaned_columns_1.append(column_name)
        df1 = pd.DataFrame(table_data_1[1:], columns=cleaned_columns_1)

        domain_col = df1.columns[0]
        std_score_col = df1.columns[2]
        percentile_col = df1.columns[3]

        selected_columns_1 = df1[[domain_col, std_score_col, percentile_col]].copy()
        selected_columns_1.columns = ["Domain Scores", "Standard Score", "Percentile"]

        selected_columns_1['Percentile'] = selected_columns_1['Percentile'].apply(clean_percentile)
        selected_columns_1['Standard Score'] = selected_columns_1['Standard Score'].apply(clean_score)
        selected_columns_1.dropna(subset=['Percentile'], inplace=True)
        df1 = selected_columns_1

    # Lower Table (Domain, Score, Severity)
    df2 = None
    for table in tables:
        if len(table) > 0 and len(table[0]) >= 3 and table[0][0] == "Domain" and table[0][1] == "Score" and table[0][2] == "Severity":
            lower_table_data = table
            cleaned_lower_table_data = [row[:3] for row in lower_table_data[1:] if len(row) >= 3]
            df2 = pd.DataFrame(cleaned_lower_table_data, columns=["Domain", "Score", "Severity"])
            df2['Score'] = df2['Score'].apply(clean_score)
            df2.dropna(subset=['Score'], inplace=True)
            break

    cognitive_test_data = extract_cognitive_tests(pdf_path)

    return {"upper_table": df1, "lower_table": df2}, cognitive_test_data


def extract_cognitive_tests(pdf_path: str) -> Dict:
    """Extract cognitive test data from PDF."""
    test_data = {
        "Verbal Memory Test (VBM)": {},
        "Visual Memory Test (VSM)": {},
        "Finger Tapping Test (FTT)": {},
        "Symbol Digit Coding (SDC)": {},
        "Stroop Test (ST)": {},
        "Shifting Attention Test (SAT)": {},
        "Continuous Performance Test (CPT)": {},
        "Perception Of Emotions Test (POET)": {},
        "Reasoning Test (RT)": {},
        "Four Part Continuous Performance Test (FPCPT)": {}
    }
    with pdfplumber.open(pdf_path) as pdf:
        text = ""
        for page in pdf.pages[:3]:
            text += page.extract_text() + "\n"
    lines = text.splitlines()
    
    def find_section_indices(section_name):
        for i, line in enumerate(lines):
            if section_name in line:
                return i
        return -1
    
    def get_value_from_line(line):
        parts = line.strip().split()
        for part in reversed(parts):
            if part.isdigit():
                return part
        return None

    # Extract test data for each test type
    test_configs = [
        ("Verbal Memory Test (VBM)", [
            "Correct Hits - Immediate", "Correct Passes - Immediate",
            "Correct Hits - Delay", "Correct Passes - Delay"
        ]),
        ("Visual Memory Test (VSM)", [
            "Correct Hits - Immediate", "Correct Passes - Immediate",
            "Correct Hits - Delay", "Correct Passes - Delay"
        ]),
        ("Finger Tapping Test (FTT)", [
            "Right Taps Average", "Left Taps Average"
        ]),
        ("Symbol Digit Coding (SDC)", [
            "Correct Responses", "Errors*"
        ]),
        ("Stroop Test (ST)", [
            "Simple Reaction Time*", "Complex Reaction Time Correct*",
            "Stroop Reaction Time Correct*", "Stroop Commission Errors*"
        ]),
        ("Shifting Attention Test (SAT)", [
            "Correct Responses", "Errors*", "Correct Reaction Time*"
        ]),
        ("Continuous Performance Test (CPT)", [
            "Correct Responses", "Omission Errors*", "Commission Errors*", "Choice Reaction Time Correct*"
        ]),
        ("Perception Of Emotions Test (POET)", [
            "Correct Responses", "Average Correct Reaction Time*", "Omission Errors*", "Commission Errors*"
        ]),
        ("Reasoning Test (RT)", [
            "Correct Responses", "Average Correct Reaction Time*", "Commission Errors*", "Omission Errors*"
        ]),
    ]
    
    for test_name, metrics in test_configs:
        idx = find_section_indices(test_name)
        if idx != -1:
            for offset, metric in enumerate(metrics):
                if idx + offset + 1 < len(lines):
                    line = lines[idx + offset + 1]
                    value = get_value_from_line(line)
                    if value:
                        test_data[test_name][metric] = value
    
    # Handle POET emotions separately
    for emo in ["Positive Emotions", "Negative Emotions"]:
        emo_idx = find_section_indices(emo)
        if emo_idx != -1:
            for sub_offset, sub_metric in enumerate(["Correct Hits", "Reaction Time*"]):
                if emo_idx + sub_offset + 1 < len(lines):
                    line = lines[emo_idx + sub_offset + 1]
                    value = get_value_from_line(line)
                    if value:
                        test_data["Perception Of Emotions Test (POET)"][f"{emo} {sub_metric}"] = value
    
    # Handle FPCPT parts
    idx = find_section_indices("Four Part Continuous Performance Test")
    if idx != -1:
        for part in range(1, 5):
            part_idx = find_section_indices(f"Part {part}")
            if part_idx != -1:
                for offset, metric in enumerate([
                    "Correct Responses", "Average Correct Reaction Time*", "Incorrect Responses*", 
                    "Average Incorrect Reaction Time*", "Omission Errors*"
                ]):
                    if part_idx + offset + 1 < len(lines):
                        line = lines[part_idx + offset + 1]
                        value = get_value_from_line(line)
                        if value:
                            test_data["Four Part Continuous Performance Test (FPCPT)"][f"Part {part} {metric}"] = value
    
    return test_data


def extract_standard_scores_from_table(table: List[List]) -> List[Tuple[str, str]]:
    """Extract standard scores from a table."""
    header = table[0]
    try:
        domain_idx = header.index("Domain Scores")
    except ValueError:
        domain_idx = 0
    try:
        std_idx = header.index("Standard Score")
    except ValueError:
        std_idx = 2
    results = []
    skip_values = {None, '', 'Domain Scores', 'Standard', 'Score', 'None'}
    for row in table[1:]:
        if len(row) > max(domain_idx, std_idx):
            domain = str(row[domain_idx]).strip()
            std = row[std_idx]
            if (domain in skip_values or std in skip_values):
                continue
            results.append((domain, std))
    return results


def extract_test_standard_scores_from_tables(tables: List[List[List]]) -> Dict:
    """Extract test standard scores from tables."""
    test_scores = {}
    current_test = None
    current_subheading = None

    SUBHEADINGS = [name for name in CANONICAL_TESTS if name.startswith('Part') or name in ['Positive Emotions', 'Negative Emotions']]

    for table in tables:
        if not table or len(table) < 2:
            continue
        header = table[0]
        first_row = table[0]
        if first_row and isinstance(first_row[0], str):
            row0 = first_row[0].strip()
            for test_name in CANONICAL_TESTS:
                if test_name.lower().replace('(', '').replace(')', '') == row0.lower().replace('(', '').replace(')', ''):
                    current_test = test_name
                    current_subheading = None
                    if current_test not in test_scores:
                        test_scores[current_test] = {"__main__": []}
                    break
        for row in table[1:]:
            if not row or not row[0]:
                continue
            row0 = str(row[0]).strip()
            found_test = False
            for test_name in CANONICAL_TESTS:
                if test_name.lower().replace('(', '').replace(')', '') == row0.lower().replace('(', '').replace(')', ''):
                    current_test = test_name
                    current_subheading = None
                    if current_test not in test_scores:
                        test_scores[current_test] = {"__main__": []}
                    found_test = True
                    break
            if found_test:
                continue
            found_sub = False
            for sub in SUBHEADINGS:
                if sub.lower() == row0.lower():
                    current_subheading = sub
                    if current_test:
                        if current_subheading not in test_scores[current_test]:
                            test_scores[current_test][current_subheading] = []
                    found_sub = True
                    break
            if found_sub:
                continue
            try:
                std_idx = header.index('Standard')
            except ValueError:
                std_idx = None
            try:
                perc_idx = header.index('Percentile')
            except ValueError:
                perc_idx = None
            metric_idx = 0
            if std_idx is not None and perc_idx is not None and len(row) > max(std_idx, perc_idx):
                metric = str(row[metric_idx]).strip()
                std = row[std_idx]
                perc = row[perc_idx]
                skip_values = {None, '', 'Domain Scores', 'Standard', 'Score', 'None'}
                if (metric in skip_values or std in skip_values or perc in skip_values):
                    continue
                if current_test:
                    if current_subheading and current_subheading in test_scores[current_test]:
                        test_scores[current_test][current_subheading].append((metric, std, perc))
                    else:
                        test_scores[current_test]["__main__"].append((metric, std, perc))
    return test_scores


def docx_with_cognitive_section_and_scores(
    df: pd.DataFrame,
    test_data: Dict,
    standard_scores: List[Tuple[str, str]],
    test_standard_scores: Dict,
    upper_table_df: Optional[pd.DataFrame],
    tables: List[List[List]],
    doc: Optional[Document] = None
) -> Document:
    """Create DOCX document with cognitive section and scores."""
    if doc is None:
        doc = Document()
    
    # COGNITIVE //
    p = doc.add_paragraph()
    run = p.add_run("COGNITIVE //")
    run.bold = True
    doc.add_paragraph()
    
    # Cognitive Domain Scores header
    p = doc.add_paragraph()
    run = p.add_run("Cognitive Domain Scores (Flagged where Standard Score < 80) -")
    run.bold = True
    
    # Domain scores as bullets
    if upper_table_df is not None:
        for _, row in upper_table_df.iterrows():
            domain = row.get('Domain Scores')
            std = row.get('Standard Score') if 'Standard Score' in row else row.get('Score')
            perc = row.get('Percentile')
            try:
                std_float = float(str(std).replace('-', '').strip())
                flag = " | FLAG" if std_float < 80 else ""
            except:
                flag = ""
            p = doc.add_paragraph(style='ListBullet')
            text = f"{domain}: Standard Score: {std} | Percentile: {perc}"
            if "FLAG" in flag:
                run = p.add_run(text)
                run = p.add_run(" | FLAG")
                run.bold = True
            else:
                run = p.add_run(text + flag)
    
    doc.add_paragraph()
    
    # Specific Test Scores header
    p = doc.add_paragraph()
    run = p.add_run("Specific Test Scores (Flagged where Standard Score < 80):")
    run.bold = True
    
    # Test scores
    for test_name in CANONICAL_TESTS:
        if test_standard_scores.get(test_name):
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run(f"{test_name}:")
            run.bold = True
            test_dict = test_standard_scores[test_name]
            if "__main__" in test_dict and test_dict["__main__"]:
                for metric, std, perc in test_dict["__main__"]:
                    try:
                        std_float = float(str(std).replace('-', '').strip())
                        flag = " | FLAG" if std_float < 80 else ""
                    except:
                        flag = ""
                    p = doc.add_paragraph(style='ListBullet')
                    text = f"{metric}: Standard Score: {std} | Percentile: {perc}"
                    if "FLAG" in flag:
                        run = p.add_run(text)
                        run = p.add_run(" | FLAG")
                        run.bold = True
                    else:
                        run = p.add_run(text + flag)
            for subheading in test_dict:
                if subheading == "__main__":
                    continue
                metrics = test_dict[subheading]
                if metrics:
                    p = doc.add_paragraph()
                    run = p.add_run(subheading)
                    run.bold = True
                    for metric, std, perc in metrics:
                        try:
                            std_float = float(str(std).replace('-', '').strip())
                            flag = " | FLAG" if std_float < 80 else ""
                        except:
                            flag = ""
                        p = doc.add_paragraph(style='ListBullet')
                        text = f"{metric}: Standard Score: {std} | Percentile: {perc}"
                        if "FLAG" in flag:
                            run = p.add_run(text)
                            run = p.add_run(" | FLAG")
                            run.bold = True
                        else:
                            run = p.add_run(text + flag)
    return doc

