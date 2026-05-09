"""
Limbic Test Processing Service - Complete Version
Handles identification and extraction of limbic test data from PDFs with enhanced NPQ formatting.
"""
import re
from typing import List, Dict, Optional
import PyPDF2
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import string
import logging

try:
    from .formatter import get_formatter
    FORMATTER_AVAILABLE = True
except ImportError:
    FORMATTER_AVAILABLE = False
    get_formatter = None  # type: ignore[assignment]
    logging.warning("Formatter service not available. LLM formatting disabled.")

logger = logging.getLogger(__name__)

RESPONSE_PHRASES = [
    "More than half the days",
    "Nearly every day",
    "Several days",
    "Not at all",
    "A moderate problem",
    "A severe problem",
    "A mild problem",
    "A little bit",
    "Quite a bit",
    "Extremely",
    "Mild",
    "Moderate",
    "Moderately",
    "None",
    "Severe",
    "Once or twice a week",
    "Three or more times a week",
    "Less than once a week",
    "Never",
    "Rarely",
    "Sometimes",
    "Not a problem",
    "Often",
    "Always",
]
# Sort by length to match the longest phrase first
RESPONSE_PHRASES.sort(key=len, reverse=True)


def parse_domain_scores(text: str) -> List[Dict[str, str]]:
    """
    Parse domain scores from NPQ format text.
    
    Handles patterns like:
    "Attention 300 Severe Impulsive 180 Moderate Memory 300 Severe..."
    
    Returns:
        List of dictionaries with keys: domain, score, severity, description
    """
    domains = []
    
    # Known domain names for NPQ tests
    domain_names = [
        "Attention", "Impulsive", "Memory", "Anxiety", "Panic",
        "Depression", "Mood Stability", "Aggression", "Fatigue",
        "Sleep", "Suicide", "Pain", "Physical", "Cognitive",
        "Emotional", "Social", "Motor"
    ]
    
    # Known severity levels
    severity_levels = [
        "Severe", "Moderate", "Moderately Severe", "Mild", 
        "Minimal", "None", "Not a problem"
    ]
    
    remaining_text = text.strip()
    
    while remaining_text:
        matched = False
        
        # Try to match each known domain name
        for domain in domain_names:
            # Pattern: Domain name followed by a number (score)
            pattern = rf'\b{re.escape(domain)}\s+(\d+)\s+'
            match = re.search(pattern, remaining_text, re.IGNORECASE)
            
            if match and match.start() == 0:  # Must be at the start
                matched = True
                domain_name = domain
                score = match.group(1)
                
                # Find the position after the score
                after_score_pos = match.end()
                after_score = remaining_text[after_score_pos:].strip()
                
                # Try to match severity level
                severity = None
                severity_end_pos = 0
                
                for sev in severity_levels:
                    if after_score.lower().startswith(sev.lower()):
                        severity = sev
                        severity_end_pos = len(sev)
                        break
                
                if not severity:
                    severity = ""
                
                # Check if there's a description after severity
                description = ""
                if severity_end_pos > 0:
                    after_severity = after_score[severity_end_pos:].strip()
                    
                    # Find where the next domain starts
                    next_domain_pos = len(after_severity)
                    for next_domain in domain_names:
                        next_pattern = rf'\b{re.escape(next_domain)}\s+\d+'
                        next_match = re.search(next_pattern, after_severity, re.IGNORECASE)
                        if next_match:
                            next_domain_pos = min(next_domain_pos, next_match.start())
                    
                    description = after_severity[:next_domain_pos].strip()
                
                # Add the parsed domain
                domains.append({
                    "domain": domain_name,
                    "score": score,
                    "severity": severity,
                    "description": description
                })
                
                # Move to the remaining text after this domain
                consumed_length = match.end() + severity_end_pos + len(description)
                remaining_text = remaining_text[consumed_length:].strip()
                break
        
        if not matched:
            break
    
    return domains


def format_domain_scores_as_table(doc: Document, domains: List[Dict[str, str]], 
                                   title: str = "Domain Scores") -> None:
    """
    Format domain scores as a clean table in the Word document.
    
    Args:
        doc: Document object
        domains: List of domain dictionaries
        title: Optional title for the table
    """
    if not domains:
        return
    
    # Add a heading for the domain scores
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    
    # Determine if we have descriptions
    has_description = any(d.get("description") for d in domains)
    
    # Create table
    num_cols = 4 if has_description else 3
    table = doc.add_table(rows=1, cols=num_cols)
    table.style = 'Light Grid Accent 1'
    
    # Set column widths for better presentation
    if has_description:
        table.columns[0].width = Inches(1.8)  # Domain
        table.columns[1].width = Inches(0.8)  # Score
        table.columns[2].width = Inches(1.2)  # Severity
        table.columns[3].width = Inches(2.0)  # Description
    else:
        table.columns[0].width = Inches(2.0)  # Domain
        table.columns[1].width = Inches(1.0)  # Score
        table.columns[2].width = Inches(1.5)  # Severity
    
    # Header row
    header_cells = table.rows[0].cells
    headers = ['Domain', 'Score', 'Severity']
    if has_description:
        headers.append('Description')
    
    for i, header_text in enumerate(headers):
        cell = header_cells[i]
        cell.text = header_text
        # Make header bold
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(11)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Data rows
    for domain_data in domains:
        row_cells = table.add_row().cells
        row_cells[0].text = domain_data.get('domain', '')
        row_cells[1].text = domain_data.get('score', '')
        row_cells[2].text = domain_data.get('severity', '')
        
        if has_description:
            row_cells[3].text = domain_data.get('description', '')
        
        # Format cells
        for i, cell in enumerate(row_cells):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                # Center align score and severity columns
                if i in [1, 2]:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add spacing after table
    doc.add_paragraph()


def format_domain_scores_as_bullets(doc: Document, domains: List[Dict[str, str]],
                                      title: str = "Domain Scores") -> None:
    """
    Format domain scores as bulleted list with tab alignment.
    
    Args:
        doc: Document object
        domains: List of domain dictionaries
        title: Optional title for the section
    """
    if not domains:
        return
    
    # Add a heading
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    
    # Add each domain as a formatted bullet
    for domain_data in domains:
        domain = domain_data.get('domain', '')
        score = domain_data.get('score', '')
        severity = domain_data.get('severity', '')
        description = domain_data.get('description', '')
        
        # Create bullet paragraph
        p = doc.add_paragraph()
        
        # Format: Domain: [tab] Score - Severity [Description]
        domain_run = p.add_run(f"• {domain}:")
        domain_run.font.size = Pt(11)
        
        # Add tab
        p.add_run('\t')
        
        # Add score and severity
        value_text = f"{score} - {severity}"
        if description:
            value_text += f" ({description})"
        
        value_run = p.add_run(value_text)
        value_run.font.size = Pt(11)
        
        # Set tab stop for alignment
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(2.5))
        
        # Set paragraph formatting
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)


def detect_and_format_npq_scores(section_text: str, doc: Document) -> str:
    """
    Detect NPQ domain scores in section text and format them appropriately.
    Returns the section text with domain scores removed (they're added as a table/bullets).
    
    Args:
        section_text: The text section to process
        doc: Document object to add formatted content to
        
    Returns:
        Section text with domain scores removed
    """
    # Pattern to detect domain score sections
    # Look for multiple consecutive "Word Number Word" patterns
    domain_score_pattern = r'((?:(?:Attention|Impulsive|Memory|Anxiety|Panic|Depression|Mood Stability|Aggression|Fatigue|Sleep|Suicide|Pain|Physical|Cognitive|Emotional|Social|Motor)\s+\d+\s+(?:Severe|Moderate|Moderately Severe|Mild|Minimal|None|Not a problem)\s*)+)'
    
    match = re.search(domain_score_pattern, section_text, re.IGNORECASE)
    
    if match:
        domain_text = match.group(1)
        domains = parse_domain_scores(domain_text)
        
        if domains:
            # Format as table (you can switch to bullets if preferred)
            format_domain_scores_as_table(doc, domains, title="Domain Scores")
            # Alternative: format_domain_scores_as_bullets(doc, domains, title="Domain Scores")
            
            # Remove the domain score text from the section
            section_text = section_text.replace(domain_text, '').strip()
            # Clean up any resulting double spaces or newlines
            section_text = re.sub(r'\n\n+', '\n\n', section_text)
            section_text = re.sub(r'  +', ' ', section_text)
    
    return section_text


def final_cleanup_psqi_format(section: str) -> str:
    """
    Final cleanup pass specifically for PSQI formatting issues.
    This ensures sub-questions are always on separate lines.
    
    Args:
        section: The section text to clean
        
    Returns:
        Cleaned section text
    """
    # Split any remaining concatenated sub-questions
    # Pattern: non-newline + space + [a-z]. + space + Capital letter
    section = re.sub(
        r'([^\n])\s+([a-z]\.)\s+([A-Z])',
        r'\1\n\2 \3',
        section
    )
    
    # Also handle the case where there's a score/response before the next sub-question
    # Example: "3 - Three or more times a week b. Wake up"
    section = re.sub(
        r'((?:\d+\s+-\s+[^\n]+?))\s+([a-z]\.)\s+([A-Z])',
        r'\1\n\2 \3',
        section
    )
    
    return section


class TestIdentifier:
    """Identifies and processes limbic tests from PDF text."""
    
    def __init__(self):
        # Dictionary of test patterns with their identifiers
        self.test_patterns = {
            "ASRS": {
                "name": "Adult ADHD Self-Report Scale (ASRS-v1.1) Symptom Checklist",
                "pattern": r"Adult ADHD Self-Report Scale.*ASRS-v1\.1",
                "short_name": "ASRS-v1.1"
            },
            "OCD-A": {
                "name": "Adult OCD Inventory (OCD-A) SF-20",
                "pattern": r"Adult OCD Inventory.*OCD-A.*SF-20",
                "short_name": "OCD-A SF-20"
            },
            "AUDIT": {
                "name": "Alcohol Use Disorders Identification Test (AUDIT) SF-10",
                "pattern": r"Alcohol Use Disorders Identification Test.*AUDIT.*SF-10",
                "short_name": "AUDIT SF-10"
            },
            "DASS-21": {
                "name": "Depression, Anxiety and Stress Scale (DASS-21) SF-21",
                "pattern": r"Depression, Anxiety and Stress Scale.*DASS.*SF-21",
                "short_name": "DASS-21"
            },
            "DASS-42": {
                "name": "Depression, Anxiety and Stress Scale (DASS-42) LF-42",
                "pattern": r"Depression, Anxiety and Stress Scale.*DASS.*LF-42",
                "short_name": "DASS-42"
            },
            "DHI": {
                "name": "Dizziness Handicap Inventory (DHI) SF-25",
                "pattern": r"Dizziness Handicap Inventory.*DHI.*SF-25",
                "short_name": "DHI SF-25"
            },
            "DAST": {
                "name": "Drug Abuse Screening Test (DAST) SF-20",
                "pattern": r"Drug Abuse Screening Test.*DAST.*SF-20",
                "short_name": "DAST SF-20"
            },
            "ESS": {
                "name": "Epworth Sleepiness Scale (ESS) SF-8",
                "pattern": r"Epworth Sleepiness Scale.*ESS.*SF-8",
                "short_name": "ESS SF-8"
            },
            "FRQ": {
                "name": "Fall Risk Questionnaire (FRQ)",
                "pattern": r"Fall Risk Questionnaire.*FRQ",
                "short_name": "FRQ"
            },
            "GAD-7": {
                "name": "Generalized Anxiety Disorder 7-item (GAD-7) Scale",
                "pattern": r"Generalized Anxiety Disorder 7-item.*GAD-7.*Scale",
                "short_name": "GAD-7"
            },
            "HIQ": {
                "name": "Head Injury Questionnaire (HIQ) LF-90",
                "pattern": r"Head Injury Questionnaire.*HIQ.*LF-90",
                "short_name": "HIQ LF-90"
            },
            "HAQ": {
                "name": "Health Assessment Questionnaire (HAQ) Disability Scale",
                "pattern": r"Health Assessment Questionnaire.*HAQ",
                "short_name": "HAQ"
            },
            "LEC-5": {
                "name": "Life Events Checklist for DSM-5 (LEC-5)",
                "pattern": r"Life Events Checklist for DSM-5.*LEC-5",
                "short_name": "LEC-5"
            },
            "MFIS": {
                "name": "Modified Fatigue Impact Scale (MFIS)",
                "pattern": r"Modified Fatigue Impact Scale.*MFIS",
                "short_name": "MFIS"
            },
            "NSI": {
                "name": "Neurobehavioral Symptom Inventory (NSI) SF-22",
                "pattern": r"Neurobehavioral Symptom Inventory.*NSI.*SF-22",
                "short_name": "NSI SF-22"
            },
            "NPQ-45": {
                "name": "NeuroPsych Questionnaire (NPQ) SF-45",
                "pattern": r"NeuroPsych Questionnaire.*NPQ.*SF-45",
                "short_name": "NPQ SF-45"
            },
            "NPQ-207": {
                "name": "NeuroPsych Questionnaire (NPQ) LF-207",
                "pattern": r"NeuroPsych Questionnaire.*NPQ.*LF-207",
                "short_name": "NPQ LF-207"
            },
            "PHQ-9": {
                "name": "Patient Health Questionnaire (PHQ-9)",
                "pattern": r"Patient Health Questionnaire.*PHQ-9",
                "short_name": "PHQ-9"
            },
            "PSQI": {
                "name": "Pittsburgh Sleep Quality Index (PSQI) SF-10",
                "pattern": r"Pittsburgh Sleep Quality Index.*PSQI.*SF-10",
                "short_name": "PSQI SF-10"
            },
            "PCL-5": {
                "name": "The PTSD Checklist (PCL-5) SF-20",
                "pattern": r"The PTSD Checklist.*PCL-5.*SF-20",
                "short_name": "PCL-5 SF-20"
            }
        }

        # Subheading patterns
        self.subheading_patterns = {
            "severity": r"(.*Severity:.*)",
            "interpretation": r"(Interpretation of.*Score:.*)",
            "syndrome": r"(.*Syndrome:.*)",
            "total_score": r"(Total Score:.*)",
            "scoring": r"(Scoring:.*)",
            "results": r"(Results:.*)",
            "assessment": r"(Assessment:.*)",
            "diagnosis": r"(Diagnosis:.*)",
            "recommendation": r"(Recommendation:.*)",
            "summary": r"(Summary:.*)"
        }

    def identify_tests(self, text: str) -> List[Dict[str, str]]:
        """
        Identify all tests present in the given text.
        
        Args:
            text: The text content to analyze
            
        Returns:
            List of dictionaries containing identified tests and their details
        """
        identified_tests = []
        
        for test_id, test_info in self.test_patterns.items():
            if re.search(test_info["pattern"], text, re.IGNORECASE):
                identified_tests.append({
                    "test_id": test_id,
                    "full_name": test_info["name"],
                    "short_name": test_info["short_name"]
                })
        
        return identified_tests

    def format_subheadings(self, text: str) -> str:
        """
        Format subheadings in the text with bold markers and clean up spacing.
        
        Args:
            text: The text to format
            
        Returns:
            Formatted text with bold subheadings and normalized spacing
        """
        formatted_text = text
        
        # CRITICAL FIX: First, merge "Interpretation of" with "Total Score:" if split
        # This is a hardcoded fix for a specific, recurring issue
        formatted_text = re.sub(
            r'Interpretation\s+of\s*\n+\s*Total\s+Score\s*:',
            'Interpretation of Total Score:',
            formatted_text,
            flags=re.IGNORECASE | re.MULTILINE
        )
        
        # Handle all possible variations with flexible spacing
        formatted_text = re.sub(
            r'Interpretation\s+of\s*\n+\s*Total\s*\n*\s*Score\s*:',
            'Interpretation of Total Score:',
            formatted_text,
            flags=re.IGNORECASE | re.MULTILINE
        )
        
        # Handle single word per line case
        formatted_text = re.sub(
            r'Interpretation\s*\n+\s*of\s*\n+\s*Total\s*\n+\s*Score\s*:',
            'Interpretation of Total Score:',
            formatted_text,
            flags=re.IGNORECASE | re.MULTILINE
        )
        
        # Fix common multi-line patterns that might be split
        # Fix "More than half the" split from "days"
        formatted_text = re.sub(
            r'More than half the\s*\n\s*days',
            'More than half the days',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix "Several" split from "days"
        formatted_text = re.sub(
            r'Several\s*\n\s*days',
            'Several days',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix "Nearly every" split from "day"
        formatted_text = re.sub(
            r'Nearly every\s*\n\s*day',
            'Nearly every day',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix "Not at" split from "all"
        formatted_text = re.sub(
            r'Not at\s*\n\s*all',
            'Not at all',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix other common response phrases that might be split
        response_phrase_patterns = [
            (r'A moderate\s*\n\s*problem', 'A moderate problem'),
            (r'A severe\s*\n\s*problem', 'A severe problem'),
            (r'A mild\s*\n\s*problem', 'A mild problem'),
            (r'A little\s*\n\s*bit', 'A little bit'),
            (r'Quite a\s*\n\s*bit', 'Quite a bit'),
            (r'Once or twice a\s*\n\s*week', 'Once or twice a week'),
            (r'Three or more times a\s*\n\s*week', 'Three or more times a week'),
            (r'Less than once a\s*\n\s*week', 'Less than once a week'),
            (r'Not a\s*\n\s*problem', 'Not a problem'),
        ]
        
        for pattern, replacement in response_phrase_patterns:
            formatted_text = re.sub(pattern, replacement, formatted_text, flags=re.IGNORECASE)
        
        # Fix instructional text that gets split (e.g., "PTSD. Examinees are instructed\nto indicate")
        # Pattern: Acronym followed by period, then instruction text split across lines
        formatted_text = re.sub(
            r'((?:PTSD|ADHD|OCD|GAD-7|PHQ-9|PCL-5)\.\s+(?:Examinees|Respondents|Participants)\s+are\s+(?:instructed|asked))\s*\n\s*(to\s+)',
            r'\1 \2',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix "to indicate how much" split
        formatted_text = re.sub(
            r'to\s+indicate\s*\n\s*how\s+much',
            'to indicate how much',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Fix scale descriptions like "from 'Not at all' to 'Extremely.' Source:\nPCL-5"
        formatted_text = re.sub(
            r'(from\s+["\'].*?["\'].*?Source:)\s*\n\s*([A-Z]+)',
            r'\1 \2',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # CRITICAL: Split PSQI sub-questions onto separate lines
        # Pattern: non-newline + space + [a-z]. + space + Capital letter
        formatted_text = re.sub(
            r'([^\n])\s+([a-z]\.)\s+([A-Z])',
            r'\1\n\2 \3',
            formatted_text
        )
        
        # Now format it with bold markers - at this point it should already be on one line
        formatted_text = re.sub(
            r'Interpretation\s+of\s+Total\s+Score\s*:',
            '**Interpretation of Total Score:**',
            formatted_text,
            flags=re.IGNORECASE
        )
        
        # Format all other subheadings
        for pattern in self.subheading_patterns.values():
            matches = re.finditer(pattern, formatted_text, re.IGNORECASE)
            for match in matches:
                subheading = match.group(1)
                if not subheading.startswith('**'):
                    formatted_subheading = f"\n**{subheading}**\n"
                    formatted_text = formatted_text.replace(subheading, formatted_subheading)
        
        # Clean up excessive whitespace and normalize spacing
        # Remove multiple consecutive newlines (keep max 2)
        formatted_text = re.sub(r'\n{3,}', '\n\n', formatted_text)
        
        # Remove spaces at the beginning of lines
        formatted_text = re.sub(r'\n +', '\n', formatted_text)
        
        # Remove multiple spaces in a row (but preserve single spaces)
        formatted_text = re.sub(r'  +', ' ', formatted_text)
        
        # Ensure space between text and numbers (e.g., "Pain275" -> "Pain 275")
        formatted_text = re.sub(r'([a-zA-Z])(\d)', r'\1 \2', formatted_text)
        
        # Ensure space between numbers and text (e.g., "275Severe" -> "275 Severe")
        formatted_text = re.sub(r'(\d)([a-zA-Z])', r'\1 \2', formatted_text)
        
        # Ensure space after periods followed by letters (sentence boundaries)
        formatted_text = re.sub(r'\.([A-Z])', r'. \1', formatted_text)
        
        # Ensure space after commas if missing
        formatted_text = re.sub(r',([^\s])', r', \1', formatted_text)
        
        # Remove space before punctuation
        formatted_text = re.sub(r'\s+([.,!?:;])', r'\1', formatted_text)
        
        # Ensure test names starting with "The" are on their own line
        # This prevents incorrect merging like "SevereThe Neuropsych"
        formatted_text = re.sub(r'([a-zA-Z])The\s+([A-Z])', r'\1\n\nThe \2', formatted_text)
        
        # Remove standalone closing parentheses on their own lines
        formatted_text = re.sub(r'\n\s*\)\s*\n', '\n\n', formatted_text)
        formatted_text = re.sub(r'^\s*\)\s*$', '', formatted_text, flags=re.MULTILINE)
        
        # Use LLM formatter if available (after regex fixes)
        if FORMATTER_AVAILABLE:
            try:
                formatter = get_formatter()
                if formatter.llm is not None:  # Only use if LLM is properly initialized
                    logger.info("Applying LLM formatting to text section")
                    formatted_text = formatter.format_text(formatted_text)
            except Exception as e:
                logger.warning(f"LLM formatting failed, using regex-only formatting: {str(e)}")
        
        return formatted_text

    def get_test_sections(self, text: str) -> Dict[str, str]:
        """
        Extract the sections of text containing each identified test, excluding header information.
        """
        test_sections = {}
        identified_tests = self.identify_tests(text)
        
        # Skip header information
        header_patterns = [
            r"Test Date:.*\n",
            r"Age:.*\n",
            r"Total Test Time:.*\n",
            r"Language:.*\n",
            r"Duration:.*\n",
            r"CNSVS Online Version.*\n",
            r"Patient ID:.*\n",
            r"Administrator:.*\n"
        ]
        
        # Remove header information from text
        cleaned_text = text
        for pattern in header_patterns:
            cleaned_text = re.sub(pattern, "", cleaned_text, flags=re.IGNORECASE)
        
        for test in identified_tests:
            test_info = self.test_patterns[test["test_id"]]
            match = re.search(test_info["pattern"], cleaned_text, re.IGNORECASE)
            
            if match:
                start_pos = match.start()
                # Find the next test or end of text
                next_test_pos = len(cleaned_text)
                for other_test_id, other_test_info in self.test_patterns.items():
                    if other_test_id != test["test_id"]:
                        other_match = re.search(other_test_info["pattern"], cleaned_text[start_pos + 1:])
                        if other_match:
                            next_test_pos = min(next_test_pos, start_pos + 1 + other_match.start())
                
                # Get the section and format subheadings
                section = cleaned_text[start_pos:next_test_pos].strip()
                section = re.sub(test_info["pattern"], "", section, flags=re.IGNORECASE).strip()
                
                # CRITICAL: Hardcoded fix for "Interpretation of" + "Total Score:" split
                section = re.sub(
                    r'Interpretation\s+of\s*\n+\s*Total\s+Score\s*:',
                    'Interpretation of Total Score:',
                    section,
                    flags=re.IGNORECASE | re.MULTILINE
                )
                
                # FIX: Split PSQI sub-questions onto separate lines
                # Pattern: non-newline + space + [a-z]. + space + Capital letter
                section = re.sub(
                    r'([^\n])\s+([a-z]\.)\s+([A-Z])',
                    r'\1\n\2 \3',
                    section
                )
                
                # Remove standalone closing parentheses
                section = re.sub(r'\n\s*\)\s*\n', '\n\n', section)
                section = re.sub(r'^\s*\)\s*$', '', section, flags=re.MULTILINE)
                
                formatted_section = self.format_subheadings(section)
                
                # CRITICAL: Apply final PSQI cleanup
                if test["test_id"] == "PSQI":
                    formatted_section = final_cleanup_psqi_format(formatted_section)
                
                test_sections[test["test_id"]] = formatted_section
        
        return test_sections

    def _add_bullet_with_alignment(self, doc: Document, text: str, tab_position: float = 3.5):
        """
        Add a bullet point with consistent tab stop alignment for values.
        All values will align vertically in a column starting from the same position.
        
        Args:
            doc: Document object
            text: Text to add as bullet (should contain tab character for alignment)
            tab_position: Tab stop position in inches (default: 3.5 inches)
        """
        # Create paragraph without style to have full control over formatting
        p = doc.add_paragraph()
        
        # Split text at tab character to separate question and value
        if '\t' in text:
            parts = text.split('\t', 1)
            question_part = parts[0]
            value_part = parts[1]
            
            # Add bullet character and question part
            bullet_text = f"• {question_part}"
            question_run = p.add_run(bullet_text)
            question_run.font.size = Pt(12)
            
            # Add tab character (this will align to our tab stop)
            tab_run = p.add_run('\t')
            
            # Add value part
            value_run = p.add_run(value_part)
            value_run.font.size = Pt(12)
        else:
            # No tab, just add bullet and text
            bullet_text = f"• {text}"
            text_run = p.add_run(bullet_text)
            text_run.font.size = Pt(12)
        
        # Set consistent tab stop for vertical alignment
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(tab_position))
        
        # Set paragraph formatting for proper bullet indentation
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(6)
        
        return p

    def create_word_document(
        self,
        identified_tests: List[Dict[str, str]],
        test_sections: Dict[str, str],
        output_path: str = None,
        doc=None
    ) -> Document:
        """
        Create a Word document with formatted test sections.
        
        Args:
            identified_tests: List of identified tests
            test_sections: Dictionary of test sections
            output_path: Path to save the Word document (optional)
            doc: Optional existing Document to add content to
            
        Returns:
            The resulting Document object
        """
        if doc is None:
            doc = Document()
            
        # Add the main heading
        heading = doc.add_paragraph()
        run = heading.add_run('SCALES //')
        run.bold = True
        run.font.size = Pt(18)
        heading.paragraph_format.space_after = Pt(18)
        
        if not test_sections:
            doc.add_paragraph('No test sections were extracted.')
        else:
            for test_id, section in test_sections.items():
                # Add test name as bold heading
                test_info = self.test_patterns[test_id]
                p = doc.add_paragraph()
                run = p.add_run(test_info["name"])
                run.bold = True
                run.font.size = Pt(14)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(12)
                
                # ENHANCED: Handle NPQ tests specially - format domain scores as table
                if test_id in ["NPQ-45", "NPQ-207"]:
                    section = detect_and_format_npq_scores(section, doc)
                
                # CRITICAL FIX: One more safety check for "Interpretation of" + "Total Score:"
                # This ensures the pattern is merged before line processing
                section = re.sub(
                    r'Interpretation\s+of\s*\n+\s*Total\s+Score\s*:',
                    'Interpretation of Total Score:',
                    section,
                    flags=re.IGNORECASE | re.MULTILINE
                )
                
                # FIX: Split PSQI sub-questions onto separate lines
                # Pattern: non-newline + space + [a-z]. + space + Capital letter
                section = re.sub(
                    r'([^\n])\s+([a-z]\.)\s+([A-Z])',
                    r'\1\n\2 \3',
                    section
                )
                
                # Remove standalone closing parentheses
                section = re.sub(r'\n\s*\)\s*\n', '\n\n', section)
                section = re.sub(r'^\s*\)\s*$', '', section, flags=re.MULTILINE)
                
                # Check if this is a PSQI test
                is_psqi = test_id == "PSQI"
                
                # Preprocess lines to join wrapped lines with improved logic
                merged_lines = []
                buffer = ""
                lines = section.split('\n')
                
                for line in lines:
                    line = line.strip()
                    if not line:
                        continue
                    
                    # Skip lines that are just closing parentheses
                    if line == ')' or line == '(' or re.match(r'^\s*[\(\)]+\s*$', line):
                        continue
                    
                    parts = line.split()
                    
                    # Identify line types
                    is_new_bullet = (
                        len(parts) > 4 and
                        parts[0].isdigit() and
                        parts[1].isdigit() and
                        parts[2] == '-'
                    )
                    
                    is_heading = line.startswith('**')
                    
                    is_score_line = (
                        "score" in line.lower() or "severity" in line.lower()
                        or (len(parts) == 1 and parts[0].isdigit())
                        or (len(parts) == 2 and parts[1].isdigit() and "score" in parts[0].lower())
                        or ("score" in line.lower() and len(parts) <= 4)
                    )
                    
                    is_footnote = (
                        line.lower().startswith("developed by") or
                        line.lower().startswith("no permission required") or
                        line.lower().startswith("psychiatry research") or
                        line.lower().startswith("used with permission") or
                        line.lower().startswith("reference") or
                        line.lower().startswith("references")
                    )
                    
                    # Check if line is a citation (don't treat as new section if buffer ends with "Source:")
                    is_citation = (
                        "weathers" in line.lower() or 
                        "litz" in line.lower() or 
                        line.lower().startswith("pcl-5 (") or
                        line.lower().startswith("gad-7 (") or
                        line.lower().startswith("phq-9 (")
                    )
                    
                    # Check if line is a test name that should start new line
                    is_test_name = (
                        line.startswith("The NeuroPsych") or
                        line.startswith("The PTSD") or
                        line.startswith("The Adult") or
                        re.match(r'^The\s+[A-Z][a-z]+', line)
                    )
                    
                    # Check if line is a list item (a., b., c., etc.) - should start new line
                    is_list_item = bool(re.match(r'^[a-z]\.\s+[A-Z]', line))
                    
                    # For PSQI, check if line starts with a number (questions 1-5)
                    is_psqi_question = False
                    if is_psqi and parts and parts[0].isdigit() and len(parts[0]) <= 2:
                        is_psqi_question = True
                    
                    # Check if line starts with a response phrase fragment
                    starts_with_response_fragment = any(
                        line.lower().startswith(frag) for frag in [
                            "days", "the days", "half the days",
                            "problem", "a problem",
                            "bit", "a bit", "Total",
                            "week", "a week", "times", "time"
                        ]
                    )
                    
                    # Check if line should start a new entry (but not citations after "Source:")
                    should_start_new = is_new_bullet or is_heading or is_score_line or is_test_name or is_list_item or is_psqi_question
                    if is_footnote and not (buffer and buffer.lower().endswith("source:")):
                        should_start_new = True
                    
                    if not buffer:
                        # Start new buffer
                        buffer = line
                    elif should_start_new:
                        # Save buffer and start new
                        merged_lines.append(buffer)
                        buffer = line
                    else:
                        # Check if we should merge with previous line
                        prev_parts = buffer.split()
                        prev_is_bullet = (
                            len(prev_parts) > 4 and
                            prev_parts[0].isdigit() and
                            prev_parts[1].isdigit() and
                            prev_parts[2] == '-'
                        )
                        
                        # Check if previous line was a list item
                        prev_is_list_item = bool(re.match(r'^[a-z]\.\s+', buffer))
                        
                        # Check if previous line was a PSQI question
                        prev_is_psqi_question = False
                        if is_psqi and prev_parts and prev_parts[0].isdigit() and len(prev_parts[0]) <= 2:
                            prev_is_psqi_question = True
                        
                        buffer_lower = buffer.lower()
                        line_lower = line.lower()
                        
                        # CRITICAL: Check for "Interpretation of" + "Total Score:" pattern
                        if buffer_lower.strip().endswith("interpretation of") and line_lower.strip().startswith("total"):
                            # Force merge this specific case
                            buffer = buffer.rstrip() + " " + line.lstrip()
                            buffer = re.sub(r'\s+', ' ', buffer)
                            buffer = re.sub(r'([a-zA-Z])(\d)', r'\1 \2', buffer)
                            buffer = re.sub(r'(\d)([a-zA-Z])', r'\1 \2', buffer)
                            continue
                        
                        # Check if buffer ends with incomplete phrase or "Source:"
                        ends_with_incomplete = (
                            buffer_lower.endswith("more than half the") or
                            buffer_lower.endswith("interpretation of") or
                            buffer_lower.endswith("not at") or
                            buffer_lower.endswith("nearly every") or
                            buffer_lower.endswith("a moderate") or
                            buffer_lower.endswith("a severe") or
                            buffer_lower.endswith("a mild") or
                            buffer_lower.endswith("a little") or
                            buffer_lower.endswith("quite a") or
                            buffer_lower.endswith("less than once a") or
                            buffer_lower.endswith("once or twice a") or
                            buffer_lower.endswith("three or more times a") or
                            buffer_lower.endswith("not a") or
                            buffer_lower.endswith("source:") or
                            buffer_lower.endswith("source :") or
                            buffer_lower.endswith(" of") or
                            buffer_lower.endswith("because you...")
                        )
                        
                        # Merge if:
                        # 1. Previous line is a bullet and current is not a new structure
                        # 2. Current line starts with lowercase or response fragment (but not if it's a list item)
                        # 3. Previous line ends with incomplete phrase (including "Source:")
                        # 4. Current line is citation and buffer mentions source
                        # 5. Current line doesn't start with a number
                        # 6. Previous was list item/PSQI question and current is continuation
                        should_merge = (
                            (prev_is_bullet and not (parts and parts[0].isdigit()) and not is_list_item) or
                            (starts_with_response_fragment and not is_list_item) or
                            ends_with_incomplete or
                            (is_citation and "source" in buffer_lower) or
                            (prev_is_list_item and not is_list_item and not is_heading and not is_score_line and not is_footnote) or
                            (prev_is_psqi_question and not is_psqi_question and not is_list_item and not is_heading and not is_score_line and not is_footnote)
                        )
                        
                        if should_merge:
                            # Merge with proper spacing
                            buffer = buffer.rstrip() + " " + line.lstrip()
                            # Clean up multiple spaces
                            buffer = re.sub(r'\s+', ' ', buffer)
                            # Ensure space between text and numbers
                            buffer = re.sub(r'([a-zA-Z])(\d)', r'\1 \2', buffer)
                            buffer = re.sub(r'(\d)([a-zA-Z])', r'\1 \2', buffer)
                        else:
                            # Start new line
                            merged_lines.append(buffer)
                            buffer = line
                
                if buffer:
                    merged_lines.append(buffer)
                    
                # Process merged lines
                for line in merged_lines:
                    if line.startswith('**'):
                        p = doc.add_paragraph()
                        run = p.add_run(line.strip('*'))
                        run.bold = True
                        run.font.size = Pt(12)
                        p.paragraph_format.space_before = Pt(12)
                        p.paragraph_format.space_after = Pt(12)
                    else:
                        parts = line.split()
                        
                        # Handle list items (a., b., c., etc.) with score and response
                        # Pattern: "a. Description text 3 - Response phrase"
                        if parts and re.match(r'^[a-z]\.$', parts[0]):
                            # Find the score (number) in the line
                            score_idx = None
                            for i in range(1, len(parts)):
                                if parts[i].isdigit():
                                    # Check if next part is a dash
                                    if i + 1 < len(parts) and parts[i + 1] == '-':
                                        score_idx = i
                                        break
                            
                            if score_idx is not None:
                                # Extract parts: description, score, response
                                description = ' '.join(parts[1:score_idx]).strip()
                                # Remove trailing period from description if present
                                description = description.rstrip('.')
                                score = parts[score_idx]
                                response = ' '.join(parts[score_idx + 2:]).strip()
                                
                                # Format with proper spacing: description (tab) score - response
                                bullet_text = f"{description}:\t{score} - {response}"
                                self._add_bullet_with_alignment(doc, bullet_text, tab_position=3.5)
                                continue
                            else:
                                # List item without score, just display the description
                                description = ' '.join(parts[1:]).strip()
                                p = doc.add_paragraph(description, style='List Bullet')
                                continue
                        
                        # For PSQI, handle numbered questions (1-5)
                        if is_psqi and parts and parts[0].isdigit() and len(parts[0]) <= 2:
                            # This is a numbered question with its answer
                            p = doc.add_paragraph(line)
                            p.paragraph_format.space_after = Pt(6)
                            continue
                        
                        if (
                            len(parts) > 4 and
                            parts[0].isdigit() and
                            parts[1].isdigit() and
                            parts[2] == '-'
                        ):
                            score = parts[1]
                            dash_index = line.find('-')
                            after_dash = line[dash_index+1:].strip()
                            response = ""
                            question = after_dash
                            
                            # Try to extract response phrase from the beginning
                            for phrase in RESPONSE_PHRASES:
                                if after_dash.startswith(phrase):
                                    response = phrase
                                    question = after_dash[len(phrase):].strip()
                                    break
                            
                            # Ensure question text is properly spaced
                            question = ' '.join(question.split())
                            # Format with proper spacing: question (tab) score - response
                            bullet_text = f"{question}:\t{score} - {response}"
                            self._add_bullet_with_alignment(doc, bullet_text, tab_position=3.5)
                            continue
                        if len(parts) > 3 and parts[0][0].isdigit():
                            score_idx = None
                            for i, part in enumerate(parts):
                                if part.isdigit() and i != 0:
                                    score_idx = i
                                    break
                            if score_idx is not None and score_idx + 2 <= len(parts):
                                if parts[score_idx + 1] == '-':
                                    question = ' '.join(parts[:score_idx])
                                    # Remove leading numbers but preserve spacing
                                    question = re.sub(r'^[\d\.\)]+\s*', '', question).strip()
                                    score = parts[score_idx]
                                    response = ' '.join(parts[score_idx + 2:])
                                    # Format with proper spacing: question (tab) score - response
                                    bullet_text = f"{question}:\t{score} - {response}"
                                    self._add_bullet_with_alignment(doc, bullet_text, tab_position=3.5)
                                    continue
                        if len(parts) > 3 and (parts[0][0].isalpha() and (parts[0].endswith('.') or parts[0].endswith(')'))):
                            double_space_index = line.find('  ')
                            if double_space_index != -1:
                                left = line[:double_space_index].strip()
                                right = line[double_space_index:].strip()
                                right_parts = right.split(' - ', 1)
                                if len(right_parts) == 2:
                                    score = right_parts[0].strip()
                                    response = right_parts[1].strip()
                                    question = left[2:].strip().lstrip(string.punctuation + " ").rstrip(string.punctuation + " ")
                                    # Format with proper spacing: question (tab) score - response
                                    bullet_text = f"{question}:\t{score} - {response}"
                                    self._add_bullet_with_alignment(doc, bullet_text, tab_position=3.5)
                                    continue
                        if line.startswith('•'):
                            content = line[1:].strip()
                            if ':' in content:
                                idx = content.rfind(':')
                                question = content[:idx].strip()
                                rest = content[idx+1:].strip()
                                # Format with proper spacing: question (tab) value
                                bullet_text = f"{question}:\t{rest}"
                                self._add_bullet_with_alignment(doc, bullet_text, tab_position=3.5)
                                continue
                            else:
                                p = doc.add_paragraph(content, style='List Bullet')
                                continue
                        p = doc.add_paragraph(line)
                        p.paragraph_format.space_after = Pt(6)
                if test_id != list(test_sections.keys())[-1]:
                    doc.add_page_break()
                    
        if output_path:
            doc.save(output_path)
        return doc


def read_pdf(file_path: str) -> str:
    """
    Read and extract text from a PDF file with improved text normalization.
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        Extracted text from the PDF with normalized spacing
    """
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    # Normalize the text to fix line break issues
                    text += normalize_pdf_text(page_text) + "\n"
            return text
    except Exception as e:
        print(f"Error reading PDF file: {str(e)}")
        return ""


def normalize_pdf_text(text: str) -> str:
    """
    Normalize PDF text to handle improper line breaks and spacing.
    
    This function fixes common PDF extraction issues where:
    - Lines are broken in the middle of words or phrases
    - Response phrases like "More than half the days" are split across lines
    - Excessive spaces are left behind
    
    Args:
        text: Raw text extracted from PDF
        
    Returns:
        Normalized text with proper spacing and line breaks
    """
    if not text:
        return ""
    
    # CRITICAL FIX: Hardcoded fix for "Interpretation of" + "Total Score:"
    # This must be done BEFORE any other processing
    text = re.sub(
        r'Interpretation\s+of\s*[\r\n]+\s*Total\s+Score\s*:',
        'Interpretation of Total Score:',
        text,
        flags=re.IGNORECASE
    )
    
    # Handle multiple newlines between words
    text = re.sub(
        r'Interpretation\s+of\s*[\r\n]+\s*Total\s*[\r\n]+\s*Score\s*:',
        'Interpretation of Total Score:',
        text,
        flags=re.IGNORECASE
    )
    
    # Handle case where each word is on separate line
    text = re.sub(
        r'Interpretation\s*[\r\n]+\s*of\s*[\r\n]+\s*Total\s*[\r\n]+\s*Score\s*:',
        'Interpretation of Total Score:',
        text,
        flags=re.IGNORECASE
    )
    
    # CRITICAL: Split PSQI sub-questions BEFORE line merging
    # Pattern: any non-newline character + space + [a-z]. + space + Capital letter
    # Example: "week a. Cannot" -> "week\na. Cannot"
    text = re.sub(
        r'([^\n])\s+([a-z]\.)\s+([A-Z])',
        r'\1\n\2 \3',
        text
    )
    
    # Also handle case where there's a number/score before the sub-question
    # Example: "3 - Three or more times a week b. Wake up"
    text = re.sub(
        r'((?:\d+\s+-\s+[^\n]+?))\s+([a-z]\.)\s+([A-Z])',
        r'\1\n\2 \3',
        text
    )
    
    # FIX: Remove standalone closing parentheses
    text = re.sub(r'\n\s*\)\s*\n', '\n\n', text)
    text = re.sub(r'^\s*\)\s*$', '', text, flags=re.MULTILINE)
    
    lines = text.split('\n')
    normalized_lines = []
    i = 0
    
    # Phrases that indicate incomplete lines that should be merged
    incomplete_endings = [
        "interpretation of", "interpretation of ", 
        "more than half the", "not at", "several",
        "nearly every", "a moderate", "a severe", "a mild", "a little",
        "quite a", "once or twice a", "three or more times a", "The",
        "less than once a", "not a", "more than", "less than",
        "examinees are instructed", "respondents are instructed",
        "to indicate how much", "instructed to", "are instructed",
        " of", "of", "because you..."
    ]
    
    # Fragments that indicate continuation lines
    continuation_starters = [
        "days", "the days", "half the days", "problem", "a problem", "Neuropsych",
        "bit", "a bit", "week", "a week", "times", "time", 
        "total", "total score", "Total", "Total Score", "Total Score:",
        "score", "score:", "Score", "Score:",
        "drug", "(worse)", "to ", 
        "indicate", "to indicate"
    ]
    
    # Citations and references that should merge
    citation_keywords = [
        "weathers", "litz", "keane", "palmieri", "marx", "schnurr",
        "national center", "psychiatry research", "used with permission",
        "developed by", "source:"
    ]
    
    # Strong indicators that a new section/item is starting (don't merge)
    new_section_indicators = [
        r'^\d+\s+\d+\s+-',
        r'^\d+\.',
        r'^[a-z]\.\s+[A-Z]',  # CRITICAL: This is our sub-question pattern
        r'^[A-Z][a-z]+\s+[A-Z]',
        r'^The\s+[A-Z]',
        r'^•',
        r'^\*\*',
        r'^\d+\s*$',  # Just a number on its own line
    ]
    
    while i < len(lines):
        current_line = lines[i].rstrip()
        
        # Skip empty lines
        if not current_line:
            normalized_lines.append('')
            i += 1
            continue
        
        # Build a merged line by checking multiple continuation conditions
        merged_line = current_line
        j = i + 1
        
        while j < len(lines):
            if j >= len(lines):
                break
                
            next_line = lines[j].strip()
            
            # Skip empty lines
            if not next_line:
                break
            
            # Check if next line starts a new section/item (strong indicators)
            is_new_section = any(re.match(pattern, next_line) for pattern in new_section_indicators)
            if is_new_section:
                break
            
            # Specific check for test names that should NEVER be merged
            next_is_test_name = (
                next_line.startswith("The NeuroPsych") or
                next_line.startswith("The PTSD") or
                next_line.startswith("The Adult") or
                re.match(r'^The\s+[A-Z][a-z]+', next_line)
            )
            if next_is_test_name:
                break
            
            # CRITICAL: Check for list items (a., b., c., etc.) - should NEVER be merged
            next_is_list_item = re.match(r'^[a-z]\.\s+[A-Z]', next_line)
            if next_is_list_item:
                break
            
            should_merge = False
            current_lower = merged_line.lower().rstrip()
            next_lower = next_line.lower()
            
            # Case 1: Current line ends with incomplete phrase
            ends_with_incomplete = any(
                current_lower.endswith(phrase.lower()) 
                for phrase in incomplete_endings
            )
            if ends_with_incomplete:
                should_merge = True
            
            # Case 2: Next line starts with continuation fragment
            # BUT NOT if next line is a sub-question
            starts_with_continuation = any(
                next_lower.startswith(frag.lower()) 
                for frag in continuation_starters
            )
            if starts_with_continuation and not next_is_list_item:
                should_merge = True
            
            # CRITICAL: Special case for "Interpretation of" + "Total Score:"
            if "interpretation of" in current_lower and ("total" in next_lower or "score" in next_lower):
                should_merge = True
            
            if current_lower.strip().endswith("interpretation of") and next_lower.startswith("total"):
                should_merge = True
            
            # Case 3: Current line ends with "Source:" or citation context
            ends_with_source = current_lower.endswith("source:") or current_lower.endswith("source :")
            if ends_with_source:
                should_merge = True
            
            # Case 4: Next line contains citation info and previous mentions source/reference
            next_is_citation = any(keyword in next_lower for keyword in citation_keywords)
            prev_mentions_source = any(
                keyword in current_lower 
                for keyword in ["source", "reference", "developed by", "used with"]
            )
            if next_is_citation and prev_mentions_source:
                should_merge = True
            
            # Case 5: Current line ends with period after acronym
            if merged_line.endswith('.'):
                words = merged_line.split()
                if words:
                    last_word = words[-1].rstrip('.')
                    is_acronym = last_word.isupper() and 2 <= len(last_word) <= 6
                    next_starts_sentence = next_line and next_line[0].isupper()
                    instructional_starts = any(
                        next_lower.startswith(phrase) 
                        for phrase in ["examinees", "respondents", "participants", "the"]
                    )
                    if is_acronym and next_starts_sentence and instructional_starts:
                        should_merge = True
            
            # Case 6: Next line starts with lowercase (but not if it's a sub-question)
            next_starts_lowercase = next_line and next_line[0].islower() and not next_is_list_item
            current_not_sentence_end = merged_line and merged_line[-1] not in '.!?'
            if next_starts_lowercase and current_not_sentence_end:
                should_merge = True
            
            # Case 7: Current line seems incomplete
            seems_incomplete = (
                merged_line and 
                merged_line[-1] not in '.!?' and
                not merged_line.endswith(':') and
                next_line and
                not next_line[0].isdigit() and
                not next_is_list_item
            )
            next_starts_new_sentence = (
                next_line and 
                len(next_line.split()) > 0 and
                next_line[0].isupper() and
                next_line.split()[0] not in ["Total", "Score", "Interpretation"]
            )
            if seems_incomplete and not next_starts_new_sentence:
                continuation_words = ['and', 'or', 'but', 'the', 'a', 'an', 'of', 'in', 'on', 'at', 'by', 'for', 'with', 'from', 'to']
                if any(next_lower.startswith(word + ' ') or next_lower == word for word in continuation_words):
                    should_merge = True
            
            if should_merge:
                merged_line = merged_line.rstrip() + " " + next_line.lstrip()
                j += 1
            else:
                break
        
        # Clean up multiple spaces but preserve single spaces
        merged_line = re.sub(r'  +', ' ', merged_line).strip()
        
        # Ensure proper spacing around punctuation
        merged_line = re.sub(r'\s+([.,!?:;])', r'\1', merged_line)
        merged_line = re.sub(r'([.,!?:;])([A-Za-z])', r'\1 \2', merged_line)
        
        # Add space between text and numbers
        merged_line = re.sub(r'([a-zA-Z])(\d)', r'\1 \2', merged_line)
        
        # Add space between numbers and text
        merged_line = re.sub(r'(\d)([a-zA-Z])', r'\1 \2', merged_line)
        
        normalized_lines.append(merged_line)
        
        # Move to next unprocessed line
        i = j if j > i else i + 1
    
    result = '\n'.join(normalized_lines)
    
    # Final cleanup: ensure consistent spacing
    result = re.sub(r'  +', ' ', result)
    result = re.sub(r'([a-zA-Z])(\d)', r'\1 \2', result)
    result = re.sub(r'(\d)([a-zA-Z])', r'\1 \2', result)
    
    # CRITICAL FINAL STEP: One more pass to split any remaining concatenated sub-questions
    # This catches any that might have been created during merging
    result = re.sub(
        r'([^\n])\s+([a-z]\.)\s+([A-Z])',
        r'\1\n\2 \3',
        result
    )
    
    # Remove standalone closing parentheses
    result = re.sub(r'\n\s*\)\s*\n', '\n\n', result)
    result = re.sub(r'^\s*\)\s*$', '', result, flags=re.MULTILINE)
    
    # Remove standalone opening parentheses
    result = re.sub(r'\n\s*\(\s*\n', '\n\n', result)
    result = re.sub(r'^\s*\(\s*$', '', result, flags=re.MULTILINE)
    
    return result